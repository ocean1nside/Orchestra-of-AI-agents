from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument

from orchestrator_api.core.settings import get_settings


def _read_docx_text(raw: bytes) -> str:
    doc = DocxDocument(BytesIO(raw))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = (cell.text or "").strip()
                if t:
                    parts.append(t)
    return "\n\n".join(parts)


def extract_text_from_bytes(*, filename: str, data: bytes) -> tuple[str, str]:
    """Извлекает текст из загруженных байт (.md, .txt, .docx)."""
    if not data:
        raise ValueError("Empty file.")
    suffix = Path(filename).suffix.lower()

    if suffix == ".docx":
        text = _read_docx_text(data)
        if not text.strip():
            raise ValueError("Empty document after DOCX extraction")
        return text, suffix

    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as e:
        raise ValueError("Only UTF-8 text files are supported for .md/.txt") from e

    if suffix not in {".md", ".txt", ""}:
        raise ValueError(f"Unsupported file type: {suffix or 'no suffix'}")

    return text, suffix or ".txt"


def read_original_text(*, relative_path: str) -> tuple[str, str]:
    """
    Извлекает текст из оригинала в storage. Возвращает (text, suffix).
    Поддерживается: .md, .txt, .docx (UTF-8 для md/txt).
    """
    root = Path(get_settings().storage_path)
    full = (root / relative_path).resolve()
    if not str(full).startswith(str(root.resolve())):
        raise ValueError("Invalid original path")
    if not full.is_file():
        raise FileNotFoundError(f"Original not found: {relative_path}")

    return extract_text_from_bytes(filename=full.name, data=full.read_bytes())
